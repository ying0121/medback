"""
Convert Medical-Bot-Console-User-Manual.html into a Word (.docx) file.
"""
from __future__ import annotations

import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
HTML_PATH = ROOT / "Medical-Bot-Console-User-Manual.html"
OUT_PATH = ROOT / "Medical-Bot-Console-User-Manual.docx"
IMAGES = ROOT / "images"

BRAND = RGBColor(0x0A, 0x53, 0x68)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)


def set_run_font(run, *, size=11, bold=False, color=INK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0A5368")
    pBdr.append(bottom)
    pPr.append(pBdr)


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            p.remove(child)


def append_inline(paragraph, node, *, base_size=11, base_bold=False, color=INK):
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        run = paragraph.add_run(text)
        set_run_font(run, size=base_size, bold=base_bold, color=color)
        return

    if not isinstance(node, Tag):
        return

    name = node.name.lower()
    if name in ("script", "style"):
        return

    if name == "br":
        paragraph.add_run().add_break()
        return

    if name in ("strong", "b"):
        for child in node.children:
            append_inline(paragraph, child, base_size=base_size, base_bold=True, color=color)
        return

    if name in ("em", "i"):
        for child in node.children:
            append_inline(paragraph, child, base_size=base_size, base_bold=base_bold, color=color)
            if paragraph.runs:
                paragraph.runs[-1].italic = True
        return

    if name == "span" and "kbd" in (node.get("class") or []):
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=base_size - 1, bold=True, color=BRAND, name="Consolas")
        return

    if name == "a":
        for child in node.children:
            append_inline(paragraph, child, base_size=base_size, base_bold=base_bold, color=BRAND)
        return

    for child in node.children:
        append_inline(paragraph, child, base_size=base_size, base_bold=base_bold, color=color)


def add_rich_paragraph(doc, element, *, style=None, size=11, bold=False, color=INK, space_after=8):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for child in element.children:
        append_inline(p, child, base_size=size, base_bold=bold, color=color)
    if not p.runs and element.get_text(strip=True):
        run = p.add_run(element.get_text(" ", strip=True))
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def resolve_image(src: str) -> Path | None:
    if not src:
        return None
    path = (ROOT / src).resolve()
    return path if path.exists() else None


def add_image(doc, src: str, *, width_in: float, caption: str | None = None):
    path = resolve_image(src)
    if not path:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(12)
        run = cap.add_run(caption)
        set_run_font(run, size=9, color=MUTED)


def add_table(doc, table_tag: Tag):
    rows = table_tag.find_all("tr")
    if not rows:
        return
    cols = max(len(r.find_all(["th", "td"])) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, tr in enumerate(rows):
        cells = tr.find_all(["th", "td"])
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell_tag = cells[c_idx] if c_idx < len(cells) else None
            # clear default paragraph
            cell.text = ""
            p = cell.paragraphs[0]
            if cell_tag is None:
                continue
            is_header = cell_tag.name == "th" or r_idx == 0 and tr.find("th")
            for child in cell_tag.children:
                append_inline(
                    p,
                    child,
                    base_size=9 if is_header else 10,
                    base_bold=bool(is_header),
                    color=BRAND if is_header else INK,
                )
            if not p.runs:
                run = p.add_run(cell_tag.get_text(" ", strip=True))
                set_run_font(
                    run,
                    size=9 if is_header else 10,
                    bold=bool(is_header),
                    color=BRAND if is_header else INK,
                )
            if is_header:
                shade_cell(cell, "F1F5F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def process_block(doc, el: Tag):
    if not isinstance(el, Tag):
        return
    name = el.name.lower()
    classes = el.get("class") or []

    if name == "h2":
        p = add_rich_paragraph(doc, el, size=16, bold=True, color=BRAND, space_after=10)
        p.paragraph_format.space_before = Pt(18)
        add_horizontal_line(p)
        return

    if name == "h3":
        p = add_rich_paragraph(doc, el, size=13, bold=True, color=INK, space_after=6)
        p.paragraph_format.space_before = Pt(12)
        return

    if name == "h4":
        add_rich_paragraph(doc, el, size=11, bold=True, color=BRAND, space_after=4)
        return

    if name == "p":
        size = 11
        color = INK
        if "lead" in classes:
            size = 12
            color = MUTED
        if "subtitle" in classes:
            size = 12
            color = MUTED
        add_rich_paragraph(doc, el, size=size, color=color, space_after=8)
        return

    if name == "figure":
        img = el.find("img")
        cap = el.find("figcaption")
        if img:
            add_image(
                doc,
                img.get("src", ""),
                width_in=5.8,
                caption=cap.get_text(" ", strip=True) if cap else None,
            )
        return

    if name == "div" and "logo-showcase" in classes:
        img = el.find("img")
        if img:
            add_image(doc, img.get("src", ""), width_in=1.2)
        meta = el.select_one(".logo-meta")
        if meta:
            strong = meta.find("strong")
            if strong:
                add_rich_paragraph(doc, strong, size=14, bold=True, color=BRAND, space_after=4)
            for p in meta.find_all("p", recursive=False):
                add_rich_paragraph(doc, p, size=11, color=MUTED, space_after=8)
        return

    if name == "div" and "callout" in classes:
        label = el.select_one(".label")
        box = doc.add_paragraph()
        box.paragraph_format.space_before = Pt(6)
        box.paragraph_format.space_after = Pt(2)
        if label:
            run = box.add_run(label.get_text(strip=True).upper() + " — ")
            set_run_font(run, size=10, bold=True, color=BRAND)
        # remaining text from paragraphs
        texts = []
        for p in el.find_all("p"):
            texts.append(p.get_text(" ", strip=True))
        run = box.add_run(" ".join(texts))
        set_run_font(run, size=10, color=INK)
        return

    if name == "div" and "card" in classes:
        for child in el.children:
            if isinstance(child, Tag):
                process_block(doc, child)
        return

    if name == "div" and "grid-2" in classes:
        for card in el.select(".card"):
            process_block(doc, card)
        return

    if name in ("ul", "ol"):
        ordered = name == "ol"
        is_steps = "steps" in classes
        for idx, li in enumerate(el.find_all("li", recursive=False), start=1):
            style = "List Number" if ordered or is_steps else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(4)
            # For custom steps, keep numbering from Word list style
            for child in li.children:
                append_inline(p, child, base_size=11, color=INK)
            if not p.runs:
                run = p.add_run(li.get_text(" ", strip=True))
                set_run_font(run, size=11, color=INK)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return

    if name == "table":
        add_table(doc, el)
        return

    if name == "nav" and "toc" in classes:
        heading = el.find("h2")
        if heading:
            p = add_rich_paragraph(doc, heading, size=12, bold=True, color=MUTED, space_after=8)
            p.paragraph_format.space_before = Pt(6)
        ol = el.find("ol")
        if ol:
            process_block(doc, ol)
        return

    if name == "div" and "cover-meta" in classes:
        for child in el.find_all("div", recursive=False):
            add_rich_paragraph(doc, child, size=10, color=MUTED, space_after=2)
        return

    if name == "div" and "footer" in classes:
        for p in el.find_all("p"):
            add_rich_paragraph(doc, p, size=9, color=MUTED, space_after=4)
        return

    # Generic container
    if name in ("div", "section", "main", "header"):
        for child in el.children:
            if isinstance(child, Tag):
                process_block(doc, child)


def build_cover(doc, cover: Tag):
    # Logo
    img = cover.select_one(".cover-brand img")
    if img:
        add_image(doc, img.get("src", ""), width_in=1.15)

    brand_text = cover.select_one(".cover-brand-text")
    if brand_text:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(brand_text.get_text(strip=True).upper())
        set_run_font(run, size=10, bold=True, color=MUTED)

    product = cover.select_one(".cover-brand div div:last-child")
    # More reliable: second text block under cover-brand
    brand_blocks = cover.select(".cover-brand > div > div")
    if len(brand_blocks) >= 2:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(brand_blocks[1].get_text(strip=True))
        set_run_font(run, size=22, bold=True, color=BRAND)

    h1 = cover.find("h1")
    if h1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(h1.get_text(strip=True))
        set_run_font(run, size=28, bold=True, color=INK)

    subtitle = cover.select_one(".subtitle")
    if subtitle:
        p = add_rich_paragraph(doc, subtitle, size=12, color=MUTED, space_after=14)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = cover.select_one(".cover-meta")
    if meta:
        for child in meta.find_all("div", recursive=False):
            p = add_rich_paragraph(doc, child, size=10, color=MUTED, space_after=2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # page break after cover
    doc.add_page_break()


def main():
    html = HTML_PATH.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "lxml")

    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    cover = soup.select_one("header.cover")
    if cover:
        build_cover(doc, cover)

    content = soup.select_one("main.content")
    if content:
        for child in content.children:
            if isinstance(child, Tag):
                process_block(doc, child)

    # Core properties
    core = doc.core_properties
    core.title = "Medical Bot Console User Manual"
    core.author = "Healthcare Chat Bot"
    core.subject = "Clinic operations manual for Medical Bot Console"

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
